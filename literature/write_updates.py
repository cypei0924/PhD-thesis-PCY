#!/usr/bin/env python3
"""Write UPDATES_2026-10.xlsx: what the October 2026 check adds to or changes in your table
SEA_peatland_WTD_datasets.xlsx, in that table's layout.

usage: python3 write_updates.py SEA_peatland_WTD_datasets.xlsx [OUT.xlsx]

Sheets:
  New rows             rows to paste into your table (A34, C60-C65, D15), in its layout and colours, and rows to add
                       to entries you already have (ID column "A19 +")
  Corrections          every cell of your table that the papers or data files change: your row number, old and new value
  Duplicates           your rows that repeat an existing entry
  Checked, not added   datasets and reports the search found and left out, with the reason
Your table is only the template and the source of the old values; it is not kept in this repository.
Needs openpyxl; the entries come from format_expanded.py and the facts from datasets_expanded.py.
"""
import os
import sys

import openpyxl

import format_expanded as fx
from datasets_expanded import CHECKED_NOT_ADDED, DUPLICATES

HERE = os.path.dirname(os.path.abspath(__file__))
NEW_IDS = ("A34", "C60", "C61", "C62", "C63", "C64", "C65", "D15")


def fill_key(cell):
    f = cell.fill
    if not f.fill_type:
        return None
    c = f.fgColor
    return (c.type, c.theme if c.type == "theme" else c.rgb, round(c.tint or 0, 2))


def read_table(path):
    """Your table: {ID: [excel rows]}, the main sheet and {fill: legend label}."""
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        if [str(ws.cell(1, c).value or "").strip() for c in range(1, len(fx.COLS) + 1)] == fx.COLS:
            break
    else:
        sys.exit("%s: no sheet with the SEA_peatland_WTD_datasets header row" % path)
    legend, groups, cur = {}, {}, None
    for r in range(2, ws.max_row + 1):
        label = str(ws.cell(r, 2).value or "").strip()
        if label.lower() in fx.LABELS and ws.cell(r, 1).value is None:
            legend[fill_key(ws.cell(r, 1))] = label.lower()
            cur = None
            continue
        v = ws.cell(r, 1).value
        if v is not None and str(v).strip():
            cur = str(v).strip()
            groups[cur] = []
        if cur and any(ws.cell(r, c).value is not None for c in range(2, len(fx.COLS) + 1)):
            groups[cur].append(r)
    return ws, groups, legend


def corrections(ws, groups, legend):
    """TABLE_FIXES checked against your table: (ID, your row, site, column, old, new, source), skipping cells that
    already hold the new value; and the 'New row' fixes as (ID, row tuple, source)."""
    col = {h: i + 1 for i, h in enumerate(fx.COLS)}
    out, add_rows = [], []
    for i, spec, column, new, src, _ in fx.fixes_for("table"):
        rows = groups.get(i)
        if not rows:
            out.append((i, "", "", column, "(ID not in your table)", new, src))
            continue
        if column == "New row":
            if not any(fx.same(ws.cell(r, col["Site / dataset"]).value, new[0]) for r in rows):
                add_rows.append((i, new, src))
                out.append((i, "", new[0], column, "", " | ".join(str(x) for x in new), src))
            continue
        texts = [(ws.cell(r, col["Site / dataset"]).value, " ".join(str(ws.cell(r, col[c]).value or "") for c in fx.MATCH[1:]))
                 for r in rows]
        ks = fx.pick_rows(texts, spec)
        if not ks:
            out.append((i, "", spec, column, "(no row matches '%s')" % spec, new, src))
            continue
        if column == "Colour" or column in fx.ENTRY_COLS:
            ks = ks[:1]
        for k in ks:
            r = rows[k] if column != "Colour" and column not in fx.ENTRY_COLS else rows[0]
            site = ws.cell(rows[k], col["Site / dataset"]).value or ws.cell(rows[k], col["Land use / condition"]).value or ""
            if column == "Colour":
                key = fill_key(ws.cell(r, 1))
                old = legend.get(key, "no fill" if key is None else "other colour")
                if old == fx.NAMES[fx.colour(new)]:
                    continue
            else:
                old = ws.cell(r, col[column]).value
                if fx.same(old, new):
                    continue
            out.append((i, r, site, column, old, new, src))
    out.sort(key=lambda f: (f[0][0], int(f[0][1:])))    # by ID; within an ID, in the order of TABLE_FIXES
    return out, add_rows


def main(table, out):
    ws, groups, legend = read_table(table)
    fixes, add_rows = corrections(ws, groups, legend)

    wb, sh, st = fx.load_template(table)
    sh.title = "New rows"
    entries = [dict(e) for e in fx.E if e["new"] in NEW_IDS]
    for i, row, src in add_rows:            # rows for entries you already have
        entries.append(dict(new=i + " +", pubs=None, links=None, data=None, cat=None, rows=[row],
                            dup="Add below your %s rows (%s)" % (i, src)))
    notes = []
    for e in entries:
        notes += [e["dup"]] + [None] * (len(e["rows"]) - 1)
    last = fx.write_rows(sh, st, entries)
    note_col = len(fx.COLS) + 1
    sh.cell(1, note_col, "Note (not a column of your table)").font = st["head"]
    for k, n in enumerate(notes):
        if n:
            sh.cell(2 + k, note_col, n).font = st["base"]
    sh.column_dimensions[openpyxl.utils.get_column_letter(note_col)].width = 45
    fx.write_legend(sh, st, last + 3)

    fx.plain_sheet(wb, st, "Corrections", ["ID", "Your row", "Site / row", "Column", "Old value (your table)", "New value", "Source"],
                   [list(f) for f in fixes], [7, 9, 34, 20, 40, 60, 40])
    fx.plain_sheet(wb, st, "Duplicates", ["Your ID", "Same as", "Note"], [list(d) for d in DUPLICATES], [9, 18, 110])
    ck = fx.plain_sheet(wb, st, "Checked, not added", ["Item", "Link", "What it is", "Why it was left out"],
                        [list(c) for c in CHECKED_NOT_ADDED], [38, 50, 55, 60])
    for row in ck.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            if str(c.value).startswith("http"):
                c.hyperlink = c.value
                c.font = st["link"]
    wb.active = 0
    wb.properties.lastModifiedBy = None
    wb.save(out)
    print("new entries %d (+%d added rows), corrections %d, duplicates %d, checked %d -> %s" % (
        len(NEW_IDS), len(add_rows), len(fixes), len(DUPLICATES), len(CHECKED_NOT_ADDED), out))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "UPDATES_2026-10.xlsx"))
