#!/usr/bin/env python3
"""Write EXPANDED_2006-2026.xlsx in the layout of SEA_peatland_WTD_datasets.xlsx (one row per site, IDs continuing
that table: A15+, B4+, C15+, D9+, and the October 2026 additions A34, C60-C65, D15). Row facts come from
build_expanded.py; the corrections in TABLE_FIXES (datasets_expanded.py), found by reading the papers and data files,
are written over them and listed in the sheet "Changes after reading papers". The sheet "ID mapping" links each new ID
to the ID used in EXPANDED_2006-2026.md.

usage: python3 format_expanded.py SEA_peatland_WTD_datasets.xlsx [OUT.xlsx]
The table is used as the template (theme, fonts, column widths, header and colour legend), so it is not kept in
this repository; only its header row and legend are carried over. write_updates.py imports the entries from here.
"""
import copy
import os
import sys

import openpyxl
import openpyxl.comments

from build_expanded import APERS_SITES, ROWS
from datasets_expanded import TABLE_FIXES

HERE = os.path.dirname(os.path.abspath(__file__))
D = "https://doi.org/"

# ---------------------------------------------------------------- entries
# rows: (site, country, province, lat, lon, coord_source, land_use, period, resolution, method, sites, wells)
# cat: can = can use, rep = replicate, req = on request, na = not accessible, None = no fill (as in the user's C/D rows)
E = []


def add(new, old, cat, pubs, links, data, rows, dup="", window="", via=""):
    E.append(dict(new=new, old=old, cat=cat, pubs=pubs, links=links, data=data, rows=rows, dup=dup, window=window, via=via))


# ---- A: public repositories
add("A15", "A15", "rep", "Tang et al. 2020, GCB; Nishina et al. 2023, STOTEN",
    D + "10.1111/gcb.15332 ; " + D + "10.1016/j.scitotenv.2023.162062", D + "10.6084/m9.figshare.25299358",
    [("Maludam NP undrained PSF (MY-MLM forest)", "Malaysia", "Sarawak (Betong)", 1.4536, 111.1495, "FLUXNET",
      "Primary peat swamp forest", "2011/01-2014/12", "Daily", "Daily mean WT (cm)", 1, ""),
     ("Naman oil palm plantation (converted)", "Malaysia", "Sarawak (Sibu, Naman plantation)", 2.186, 111.8459, "literature",
      "Oil palm plantation converted from peat swamp forest", "2018/03-2019/05", "Daily", "", 1, "")],
    "Your A4 rows 9-10 (Tang 2020 2011-2014; Nishina 2023 Naman); this adds the public data file")
add("A16", "A16", "can", "Tang et al. 2018, GRL", D + "10.1029/2017gl076457", D + "10.5281/zenodo.1161966",
    [("Maludam MY-MLM tower", "Malaysia", "Sarawak (Betong)", 1.4536, 111.1495, "FLUXNET", "Primary peat swamp forest",
      "2013/11-2013/12", "Half-hourly", "EC CH4 flux + WT (cm)", 1, "")],
    "Related to your A6 (same tower; different period and data file)")
add("A17", "A17", "can", "Warren-Thomas et al. 2022, J. Appl. Ecol.", D + "10.1111/1365-2664.14135", D + "10.5061/dryad.rr4xgxd9v",
    [("Oil palm smallholdings", "Indonesia", "Jambi", "See dataset", "See dataset", "—",
      "Smallholder oil palm (12-month mean WT −52 to −3 cm)", "2018/08-2019/08", "Logger (see dataset)", "Water level loggers", 41, ""),
     ("Protected peat swamp forest plots", "Indonesia", "Jambi", "See dataset", "See dataset", "—",
      "Protected PSF (mean WT −3 to +15 cm)", "2018/08-2019/08", "Logger (see dataset)", "", 21, "")])
add("A18", "A18", "can", "Swails et al. 2021, Front. Environ. Sci.", D + "10.3389/fenvs.2021.617828",
    D + "10.17528/CIFOR/DATA.00201 ; " + D + "10.3389/fenvs.2021.617828.s001",
    [("Undrained peat swamp forests", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Undrained PSF",
      "~1.5 years (see paper)", "Monthly", "Manual dipwells with CH4/N2O sampling", 3, ""),
     ("Oil palm plantations", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Oil palm",
      "~1.5 years (see paper)", "Monthly", "", 3, "")],
    "Probably the same 6 sites as your A14 (Swails 2019, FOR-1-3 and OP-2007/2009/2011)")
add("A19", "A19", "can", "Hergoualc'h et al. 2026 (N application); Comeau et al. 2016, Geoderma", D + "10.1016/j.geoderma.2016.01.016",
    D + "10.17528/CIFOR/DATA.00330 ; " + D + "10.17528/CIFOR/DATA.00331 ; " + D + "10.17528/CIFOR/DATA.00290 ; " + D + "10.17528/CIFOR/DATA.00316",
    [("Forest-to-oil-palm transition and N-fertilisation plots", "Indonesia", "Jambi", "See dataset", "See dataset", "—",
      "Peat swamp forest; oil palm", "See dataset", "Monthly", "Soil respiration + environmental variables (WT column not confirmed)", "See dataset", "")])
add("A20", "A20", "can", "CIFOR SWAMP datasets (no paper identified)", "—",
    D + "10.17528/CIFOR/DATA.ZORCAF ; " + D + "10.17528/CIFOR/DATA.F5DM1Y",
    [("SWAMP peatland GHG, Katingan", "Indonesia", "Central Kalimantan (Katingan)", "See dataset", "See dataset", "—",
      "Peat swamp forest", "2015", "See dataset", "GHG fluxes + site variables (WT not confirmed)", "See dataset", ""),
     ("SWAMP peatland GHG, Tanjung Puting", "Indonesia", "Central Kalimantan (Tanjung Puting)", "See dataset", "See dataset", "—",
      "Peat swamp forest", "2015", "See dataset", "", "See dataset", "")])
add("A21", "A21", "can", "Cooper et al. 2020, Nat. Commun.", D + "10.1038/s41467-020-14298-w",
    D + "10.1038/s41467-020-14298-w (Source Data + Supplementary Data 1)",
    [("Secondary peat swamp forest", "Malaysia", "Selangor (North Selangor PSF)", "See paper", "See paper", "—",
      "Secondary PSF", "2 years (see paper)", "Monthly", "Manual dipwells; WT also at each GHG sampling", 1, "2 dipwells"),
     ("Drained forest", "Malaysia", "Selangor (North Selangor PSF)", "See paper", "See paper", "—", "Drained PSF",
      "See paper", "Per GHG sampling", "", 1, ""),
     ("Young oil palm", "Malaysia", "Selangor (North Selangor PSF)", "See paper", "See paper", "—", "Young oil palm",
      "See paper", "Per GHG sampling", "", 1, ""),
     ("Mature oil palm", "Malaysia", "Selangor (North Selangor PSF)", "See paper", "See paper", "—", "Mature oil palm",
      "See paper", "Per GHG sampling", "", 1, "")])
add("A22", "A22", "can", "Somers et al. (manuscript)", "—", D + "10.4211/hs.3953b24e0238467980a226c72cfc360e",
    [("Badas drainage canal", "Brunei", "Belait (Badas)", 4.5619, 114.3367, "dataset metadata",
      "Drainage canal in degraded peat dome", "See dataset", "See dataset", "Leveloggers in canal (streamflow), not peat WTD", 1, "")],
    "Same dome as your C12")

# ---- B: networks / on request
add("B4", "B3", "req", "Sulaiman et al. 2023, Sci. Rep.", D + "10.1038/s41598-023-27393-x", "BRIN (to be released after the project)",
    [("Palangkaraya undrained forest (Takahashi site)", "Indonesia", "Central Kalimantan (Sebangau NP)", -2.321002, 113.901161,
      "literature", "Undrained PSF (same location as A3 UF)", "1993/09-2019/12", "Daily", "Pressure sensor", 1, "")],
    "Same site as your A3/A4 UF; extends the record back to 1993")

APERS = APERS_SITES
OLD = {r[0]: r for r in ROWS}


def apers_rows(source, method):
    rows = []
    for sid, place, prov, lat, lon, drain, years, src, cover in APERS:
        if src == source:
            rows.append(("%s (%s)" % (place, sid), "Indonesia", prov, lat, lon, "Apers et al. 2022 Table B1",
                         cover if cover != "Information not available" else drain, years.replace("–", "-"), "See source", method, 1, ""))
    return rows


add("B5", "B4", "req", "Apers et al. 2022, JAMES (Table B1); SATREPS project", D + "10.1029/2021ms002784",
    "http://kalimantan88.sakura.ne.jp/", apers_rows("SATREPS", "Water level logger (graphs on project website)"))
add("B6", "B5", "req", "Putra et al. 2018, IOP EES 149; Putra et al. 2019, IOP EES 284; Hikouei et al. 2023, STOTEN; Ichsan et al. 2013 (KFCP tech. paper)",
    D + "10.1088/1755-1315/149/1/012027 ; " + D + "10.1088/1755-1315/284/1/012021 ; " + D + "10.1016/j.scitotenv.2022.159701", "—",
    [("KFCP dipwell network, ex-MRP Blocks A and E", "Indonesia", "Central Kalimantan (Kapuas)", "≈-2.3 to -2.7", "≈114.2–114.6",
      "Regional approximation", "Degraded drained peat; PSF remnants", "2010/01-2013/01", "Monthly", "Manual dipwells (blow-straw)", 1,
      "300–460 dipwells + 15 staff gauges")])
add("B7", "B6", "req", "Borneo Nature Foundation (no paper)", "—",
    "https://borneonaturefoundation.org/conservation/hydrological-monitoring-for-protecting-peatlands/",
    [("Sabangau (NLPSF) hydrology transect", "Indonesia", "Central Kalimantan", "≈-2.3", "≈113.9–114.1", "Regional approximation",
      "PSF; dammed canals", "Ongoing", "Monthly", "Manual wells + dam monitoring", 1, "32 wells (14-km transect)")])
add("B8", "B7", "rep", "Hooijer et al. 2012, BG; Evans et al. 2019, Geoderma", D + "10.5194/bg-9-1053-2012 ; " + D + "10.1016/j.geoderma.2018.12.028", "—",
    [("APRIL subsidence / WT network", "Indonesia", "Riau", "≈0.3–1.0", "≈101.9–103.0", "Regional approximation",
      "Acacia plantation; PSF", "2007-present", "2-weekly to 3-monthly", "WT in subsidence tubes", 1, "218 locations (2012); 312 sites (2019)")],
    "Your C2")
add("B9", "B8", "req", "Pratama et al. 2020, IOP MSE 796; Irfan et al. 2020, J. Phys. Conf. Ser. 1568; Irfan et al. 2023, JGSE; Irfan et al. 2026, AIP Conf. Proc.",
    D + "10.1088/1757-899x/796/1/012037 ; " + D + "10.1088/1742-6596/1568/1/012028 ; " + D + "10.26599/jgse.2023.9280008 ; " + D + "10.1063/5.0337576", "—",
    [("SESAME station, Dompas village", "Indonesia", "Riau (Bengkalis)", "See paper", "See paper", "—", "Degraded peat",
      "See paper", "Sub-daily / daily", "SESAME automatic GWL sensor", 1, ""),
     ("SESAME station SR1 (Saleh River 1)", "Indonesia", "South Sumatra", "See paper", "See paper", "—", "Degraded peat / restoration",
      "See paper (≈2018-2021)", "Sub-daily / daily", "", 1, ""),
     ("SESAME station SR2 (Saleh River 2)", "Indonesia", "South Sumatra", "See paper", "See paper", "—", "Degraded peat / restoration",
      "See paper (≈2018-2021)", "Sub-daily / daily", "", 1, ""),
     ("SESAME station LR1 (Lumpur River 1)", "Indonesia", "South Sumatra", "See paper", "See paper", "—", "Degraded peat / restoration",
      "See paper (≈2018-2021)", "Sub-daily / daily", "", 1, ""),
     ("SESAME station LR2 (Lumpur River 2)", "Indonesia", "South Sumatra", "See paper", "See paper", "—", "Degraded peat / restoration",
      "See paper (≈2018-2021)", "Sub-daily / daily", "", 1, "")])
add("B10", "B9", "req", "Khampeera et al. 2018, Walailak J. Sci. Technol.", D + "10.48048/wjst.2018.2723", "—",
    [("Kuan Kreng peat swamp (Pak Phanang Fire Control Station)", "Thailand", "Nakhon Si Thammarat", "≈8.0", "≈100.2",
      "Regional approximation", "Peat swamp forest (about 2/3 degraded)", "See paper (2010, 2012 analysed)", "Daily",
      "Station water-table records", 1, "")])
add("B11", "B10", "req", "Thai et al. 2024, Sustainability", D + "10.3390/su16020620", "—",
    [("U Minh Thuong National Park", "Vietnam", "Kien Giang", "≈9.6", "≈105.1", "Regional approximation", "Melaleuca forest on peat",
      "2002-2021", "See paper", "Park groundwater monitoring records", 1, "")])
add("B12", "B11", "req", "Pahang Peatland Restoration Project (no paper)", "—", "https://pprp.my/",
    [("SE Pahang peat swamp forest restoration area", "Malaysia", "Pahang", "≈3.2–3.4", "≈103.2–103.4", "Regional approximation",
      "PSF, restoration", "See project", "See project", "Water-level monitoring (no published series)", 1, "")],
    "Related to your C14 (Pekan canal blocking)")
add("B13", "new", "na", "Apers et al. 2022, JAMES (Table B1); Putra et al. 2025, IJEMS (39 Riau stations, 2018/10-2020/12)",
    D + "10.1029/2021ms002784 ; " + D + "10.26554/ijems.2025.9.2.46-55", "https://sipalaga.brgm.go.id",
    apers_rows("SIPALAGA", "SIPALAGA automatic station"),
    "Station list for your B1 (SiPALAGA); the A9 BRG stations belong to this network")

# ---- C: sites only in papers (no fill unless clear)
add("C15", "C13", None, "Jauhiainen et al. 2008, Ecology", D + "10.1890/07-2038.1", "—",
    [("Palangkaraya DB chamber site", "Indonesia", "Central Kalimantan", -2.3381, 114.0297, "Apers et al. 2022 Table B1",
      "Deforested, burnt, drained peat; canal dammed", "2004-2007", "See paper", "WT at chamber sites before/after damming", 1, ""),
     ("Palangkaraya DF chamber site", "Indonesia", "Central Kalimantan", -2.345, 114.0367, "Apers et al. 2022 Table B1",
      "Drained PSF", "2004-2007", "See paper", "", 1, "")])
add("C16", "C14", None, "Wösten et al. 2006, IJWRD; Wösten et al. 2008, Catena; Jaenicke et al. 2010, MASGC; Jaenicke et al. 2011, JEM; Ritzema et al. 2014, Catena",
    D + "10.1080/07900620500405973 ; " + D + "10.1016/j.catena.2007.07.010 ; " + D + "10.1007/s11027-010-9214-5 ; " + D + "10.1016/j.jenvman.2010.09.029 ; " + D + "10.1016/j.catena.2013.10.009",
    "—", [("Block C / Sebangau hydrology and canal blocking (CKPP)", "Indonesia", "Central Kalimantan", "≈-2.3 to -2.5", "≈114.0–114.2",
           "Regional approximation", "Drained/burnt peat; PSF", "≈2004-2010", "See paper", "Manual dipwell transects", "See paper", "")])
add("C17", "C15", None, "Könönen et al. 2016, Wetl. Ecol. Manag.; Lampela et al. 2017, For. Ecol. Manag.",
    D + "10.1007/s11273-016-9498-7 ; " + D + "10.1016/j.foreco.2016.12.004", "—",
    [("Sebangau logged forest", "Indonesia", "Central Kalimantan", -2.3214, 113.8953, "Apers et al. 2022 Table B1", "Logged PSF",
      "2013", "See paper", "WT record at study plots", 1, ""),
     ("Sebangau restored area", "Indonesia", "Central Kalimantan", -2.3217, 114.0181, "Apers et al. 2022 Table B1",
      "Previously deforested and drained; canal blocking, ferns", "2012-2013", "See paper", "", 1, "")])
add("C18", "C16", None, "Taufik et al. 2019, Geoderma", D + "10.1016/j.geoderma.2019.04.001", "—",
    [("Upper Sebangau PSF", "Indonesia", "Central Kalimantan", -2.42, 114.10, "Apers et al. 2022 Table B1",
      "PSF with minor influence of old canals", "2000-2008", "See paper", "GWL record", 1, ""),
     ("Air Hitam PSF", "Indonesia", "Jambi", -1.497, 104.116, "Apers et al. 2022 Table B1", "Peat swamp forest", "2003-2004",
      "See paper", "", 1, "")])
add("C19", "C17", None, "Putra et al. 2021, Hydrol. Process.; Putra et al. 2023, Mires and Peat",
    D + "10.1002/hyp.14174 ; " + D + "10.19189/map.2022.omb.sta.2407", "—",
    [("Sebangau NP, Forested", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Forested peatland",
      "2019/08-2020/01", "Hourly", "Automated + manual dipwells", 1, ""),
     ("Sebangau NP, Blocked", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Drained, with ditch dams",
      "2019/08-2020/01", "Hourly", "", 1, ""),
     ("Sebangau NP, Drained", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Drained, no ditch dams",
      "2019/08-2020/01", "Hourly", "", 1, "")])
add("C20", "C18", None, "Budi Santosa et al. 2020, Jurnal Galam", D + "10.20886/glm.2020.1.1.27-40", "—",
    [("Tumbang Nusa", "Indonesia", "Central Kalimantan (Pulang Pisau)", "≈-2.35", "≈114.09", "Regional approximation",
      "PSF / degraded", "See paper", "See paper", "WT monitoring", 1, "")], "Same area as B5 site IJ-1")
add("C21", "C19", None, "Firmansyah et al. 2020, Agritech", D + "10.30595/agritech.v22i2.8000", "—",
    [("Jabiren ex-ICCTF plot", "Indonesia", "Central Kalimantan (Pulang Pisau)", "≈-2.55", "≈114.17", "Regional approximation",
      "Degraded smallholder peat (GWT −50 to −150 cm)", "7–10 months (see paper)", "Periodic", "Wells along a transect", 1, "")])
add("C22", "C20", "rep", "Sulaeman et al. 2022, IOP EES 1025", D + "10.1088/1755-1315/1025/1/012011", "—",
    [("Time-lapse camera + WT, four sites", "Indonesia", "South Sumatra; Central Kalimantan", "See paper", "See paper", "—",
      "Various", "≈1 year", "See paper", "Camera-derived WT and peat motion", 4, "")], "Your C11")
add("C23", "C21", None, "Yulianti et al. 2024, IOP EES 1421", D + "10.1088/1755-1315/1421/1/012005", "—",
    [("Forest site", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Peat swamp forest", "2022/12-2023/06",
      "2-hourly", "Water level logger + time-lapse camera", 1, ""),
     ("Burnt site", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Burnt peat", "2022/12-2023/06", "2-hourly", "", 1, "")])
add("C24", "C22", None, "Treby et al. 2026, Pedosphere; Treby et al. 2026 (SSRN preprint)",
    D + "10.1016/j.pedsph.2026.02.009 ; " + D + "10.2139/ssrn.7163970", "—",
    [("Intact area", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Intact PSF", "2024/08-2025/11", "See paper",
      "WTD + soil water content / matric potential sensors", 1, ""),
     ("Degraded area", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—", "Historically drained and burnt",
      "2024/08-2025/11", "See paper", "", 1, "")])
add("C25", "C23", None, "Suwito et al. 2022, IOP EES 1018", D + "10.1088/1755-1315/1018/1/012027", "—",
    [("Ex-MRP canal blocking (large/medium/small blocks, unblocked)", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—",
      "Degraded PSF after fire", "See paper", "See paper", "GWL wells", 4, "")])
add("C26", "C24", None, "Itoh et al. 2017, STOTEN", D + "10.1016/j.scitotenv.2017.07.132", "—",
    [("UF", "Indonesia", "Central Kalimantan", -2.32, 113.90, "literature", "Undrained PSF", "2014/02-2015/12", "Per chamber sampling",
      "GWL at chamber measurements", 1, ""),
     ("DF", "Indonesia", "Central Kalimantan", -2.35, 114.04, "literature", "Drained PSF", "2014/02-2015/12", "Per chamber sampling", "", 1, ""),
     ("DB", "Indonesia", "Central Kalimantan", -2.34, 114.04, "literature", "Drained and burnt", "2014/02-2015/12", "Per chamber sampling", "", 1, "")],
    "Same sites as your A3")
add("C27", "C25", None, "Suryatmojo et al. 2019, IOP EES 361", D + "10.1088/1755-1315/361/1/012034", "—",
    [("Kampar, primary forest", "Indonesia", "Riau (Kampar)", "See paper", "See paper", "—", "Primary PSF", "See paper", "See paper", "Station GWL", 1, ""),
     ("Kampar, ex-fire", "Indonesia", "Riau (Kampar)", "See paper", "See paper", "—", "Burnt peat", "See paper", "See paper", "", 1, ""),
     ("Kampar, mixed plantation", "Indonesia", "Riau (Kampar)", "See paper", "See paper", "—", "Community mixed plantation", "See paper", "See paper", "", 1, "")])
add("C28", "C26", None, "Maryani et al. 2020, IOP EES 533", D + "10.1088/1755-1315/533/1/012012", "—",
    [("Siak, primary PSF", "Indonesia", "Riau (Siak)", "See paper", "See paper", "—", "Primary PSF", "2018/07-2018/12", "See paper",
      "Hydrometeorological station with GWL sensor", 1, ""),
     ("Siak, burnt peatland", "Indonesia", "Riau (Siak)", "See paper", "See paper", "—", "Burnt peat", "2018/07-2018/12", "See paper", "", 1, ""),
     ("Siak, mixed plantation", "Indonesia", "Riau (Siak)", "See paper", "See paper", "—", "Mixed plantation", "2018/07-2018/12", "See paper", "", 1, "")])
add("C29", "C27", None, "Basuki et al. 2021, IOP EES 648", D + "10.1088/1755-1315/648/1/012029", "—",
    [("Dosan village", "Indonesia", "Riau (Siak)", "See paper", "See paper", "—", "Oil palm, Acacia regrowth, shrub (mean GWT −55 cm)",
      "18 months (see paper)", "See paper", "Shallow wells + subsidence poles", "31 (both villages)", "124 shallow wells + 31 poles"),
     ("Dayun village", "Indonesia", "Riau (Siak)", "≈0.64", "≈102.03", "Regional approximation (Dayun SIPALAGA station)",
      "Oil palm, Acacia regrowth, shrub (mean GWT −66 cm)", "18 months (see paper)", "See paper", "", "", "")])
add("C30", "C28", None, "Wasilul Lutfi et al. 2021, IOP EES 756", D + "10.1088/1755-1315/756/1/012028", "—",
    [("Koto Village, Gasib", "Indonesia", "Riau (Siak)", "See paper", "See paper", "—", "Oil palm", "See paper", "Monthly",
      "GWL wells + subsidence", "See paper", "")])
add("C31", "C29", None, "Safitri et al. 2024, IOP EES 1315", D + "10.1088/1755-1315/1315/1/012058", "—",
    [("Siak canal blocking transect", "Indonesia", "Riau (Siak)", "See paper", "See paper", "—", "Various land uses near blocked canals",
      "1 year (see paper)", "See paper", "Monitoring wells (mean GWL −26.7 cm near blocked canal)", "See paper", "")])
add("C32", "C30", None, "Silviana et al. 2020, IOP MSE 796; Malik et al. 2022, IOP EES 1041",
    D + "10.1088/1757-899x/796/1/012041 ; " + D + "10.1088/1755-1315/1041/1/012047", "—",
    [("Sungai Tohor, after BRG canal blocking", "Indonesia", "Riau (Kepulauan Meranti, Tebing Tinggi)", "See paper", "See paper", "—",
      "Rewetted peat (sago, smallholder)", "1 year (see paper)", "Daily", "Water level logger in monitoring well", 1, "")],
    "Same island as your C4 (Sutikno 2020)")
add("C33", "C31", None, "Lestari et al. 2022, Forests", D + "10.3390/f13040505", "—",
    [("Rewetted vs drained peat", "Indonesia", "Riau", "See paper", "See paper", "—", "Rewetted / drained", "See paper", "See paper",
      "WT with GHG fluxes", "See paper", "")])
add("C34", "C32", None, "Nardi et al. 2021, JPSL", D + "10.29244/jpsl.11.3.442-452", "—",
    [("Kampar conservation forest", "Indonesia", "Riau (Pelalawan)", "≈0.4", "≈102.8", "Regional approximation", "Intact PSF",
      "2020/01-2020/12", "Per sampling", "WT with N2O sampling", 1, "")], "Near your A7/A8 intact site")
add("C35", "C33", None, "Sutikno et al. 2026, JDMLM", D + "10.15243/jdmlm.2026.131.9163", "—",
    [("Bengkalis, drained", "Indonesia", "Riau (Bengkalis Island)", "See paper", "See paper", "—", "Drained degraded peat",
      "2023/10-2025/04", "Daily", "In situ GWL", 1, ""),
     ("Bengkalis, undrained inland", "Indonesia", "Riau (Bengkalis Island)", "See paper", "See paper", "—", "Undrained inland peat",
      "2023/10-2025/04", "Daily", "", 1, ""),
     ("Bengkalis, undrained coastal", "Indonesia", "Riau (Bengkalis Island)", "See paper", "See paper", "—", "Undrained coastal peat",
      "2023/10-2025/04", "Daily", "", 1, "")])
add("C36", "C34", None, "Jauhiainen et al. 2012, BG", D + "10.5194/bg-9-617-2012", "—",
    [("Acacia plantation CO2 transects", "Indonesia", "Riau", "See paper", "See paper", "—", "Acacia plantation", "2 years (see paper)",
      "Monthly or quarterly", "WT in perforated PVC tubes", 1, "144 locations")], "Part of the APRIL network (your C2)")
add("C37", "C35", "rep", "Deshmukh et al. 2020, GCB", D + "10.1111/gcb.15019", "—",
    [("Kampar natural forest tower", "Indonesia", "Riau", 0.3952, 102.7646, "literature", "Intact PSF", "≈2016-2018 (see paper)",
      "Half-hourly", "Solinst Levelogger 3001, ~30 m from tower", 1, ""),
     ("Kampar Acacia plantation tower", "Indonesia", "Riau", 0.51589, 102.03641, "literature", "Acacia crassicarpa plantation",
      "≈2016-2018 (see paper)", "Half-hourly", "", 1, "")], "Your A8 (already cites Deshmukh 2020); confirms the Acacia coordinate")
add("C38", "C36", None, "Fawzi et al. 2024, Heliyon", D + "10.1016/j.heliyon.2024.e26661", "—",
    [("Coconut plantation (Water Management Trinity)", "Indonesia", "Riau (eastern Sumatra coast)", "See paper", "See paper", "—",
      "Coconut plantation (annual WTD −45 to −51 cm)", "See paper", "Per CO2 session", "Manual WTD in 4-m perforated pipes", "See paper", "")])
add("C39", "C37", None, "Dariah et al. 2013, MASGC; Husnain et al. 2014, MASGC; Comeau et al. 2016, Geoderma",
    D + "10.1007/s11027-013-9515-6 ; " + D + "10.1007/s11027-014-9550-y ; " + D + "10.1016/j.geoderma.2016.01.016", "—",
    [("Oil palm / Acacia plantations (several sites)", "Indonesia", "Jambi; Riau; South Sumatra", "See paper", "See paper", "—",
      "Oil palm; Acacia", "≈2010-2013", "Per chamber sampling", "WT at chamber sampling", "See paper", "")])
add("C40", "C38", None, "Wakhid et al. 2017, STOTEN", D + "10.1016/j.scitotenv.2017.01.035", "—",
    [("Rubber plantation on peat", "Indonesia", "See paper", "See paper", "See paper", "—", "Rubber (8 years old)", "2014/12-2015/12",
      "Monthly", "GWL with soil CO2 and subsidence", 1, "")])
add("C41", "C39", None, "Khasanah & van Noordwijk 2019, MASGC", D + "10.1007/s11027-018-9803-2", "—",
    [("Smallholder mosaic, Tanjung Jabung Barat", "Indonesia", "Jambi (Tanjung Jabung Barat)", "See paper", "See paper", "—",
      "Logged forest, rubber agroforest, coconut-coffee, betel nut-coffee, oil palm", "2012/11-2015/05", "See paper",
      "GWL with subsidence", 5, "")])
add("C42", "C40", None, "Khakim et al. 2022, GES; Maryani et al. 2021, IOP EES 810",
    D + "10.24057/2071-9388-2021-137 ; " + D + "10.1088/1755-1315/810/1/012023", "—",
    [("Air Sugihan–Air Saleh peat hydrological unit", "Indonesia", "South Sumatra", "See paper", "See paper", "—", "Degraded peat",
      "2015-2018", "See paper", "Water level observations + Sentinel-1 soil moisture model", "See paper", ""),
     ("Degraded peatland revegetation plots", "Indonesia", "South Sumatra", "See paper", "See paper", "—", "Degraded peat",
      "See paper", "See paper", "Monitoring wells in 3 x 3 m boxes", "See paper", "5 wells per box")])
add("C43", "C41", None, "Astiani et al. 2018, Biodiversitas; Herawati et al. 2018, MATEC; Nusantara et al. 2023, J. Ilmu Lingkungan; Nahda et al. 2025, BIO Web Conf.",
    D + "10.13057/biodiv/d190221 ; " + D + "10.1051/matecconf/201819503016 ; " + D + "10.14710/jil.21.4.781-788 ; " + D + "10.1051/bioconf/202516703013", "—",
    [("Kubu Raya bare degraded peat", "Indonesia", "West Kalimantan (Kubu Raya)", "See paper", "See paper", "—", "Bare degraded peat",
      "See paper", "See paper", "WT with soil CO2", "See paper", ""),
     ("Wajok Hilir, blocked vs unblocked canal", "Indonesia", "West Kalimantan (Mempawah)", "See paper", "See paper", "—",
      "Drained peat with canal blocks", "See paper", "See paper", "WT wells", 2, ""),
     ("Kubu village oil palm", "Indonesia", "West Kalimantan (Kubu Raya)", "See paper", "See paper", "—", "Oil palm (GWT −4 to −39 cm)",
      "2021/08-2021/10", "Twice daily", "Piezometers at 250, 500, 750 m", 1, "3 blocks"),
     ("Limbung village burnt peat", "Indonesia", "West Kalimantan (Kubu Raya)", "See paper", "See paper", "—", "Burnt peat",
      "2023/09-2023/12", "See paper", "Monitoring wells", 1, "")])
add("C44", "C42", None, "Novita et al. 2024, STOTEN", D + "10.1016/j.scitotenv.2024.175829", "—",
    [("Drained oil palm", "Indonesia", "West Kalimantan (Mempawah, Kubu Raya)", "See paper", "See paper", "—", "Drained oil palm",
      "See paper", "See paper", "WTL with CO2 and CH4", "See paper", ""),
     ("Rewetted oil palm", "Indonesia", "West Kalimantan (Mempawah, Kubu Raya)", "See paper", "See paper", "—", "Rewetted oil palm",
      "See paper", "See paper", "", "See paper", ""),
     ("Secondary forest", "Indonesia", "West Kalimantan (Mempawah, Kubu Raya)", "See paper", "See paper", "—", "Secondary forest",
      "See paper", "See paper", "", "See paper", "")])
add("C45", "C43", None, "Wakhid et al. 2021, Mires and Peat", D + "10.19189/map.2021.omb.sta.2159", "—",
    [("Young smallholder oil palm", "Indonesia", "South Kalimantan", "See paper", "See paper", "—", "Oil palm (7 years old)",
      "2018/09-2020/03", "Monthly", "With soil CO2", 1, "")])
add("C46", "C44", None, "Handayani et al. 2010, J. Trop. Soils", D + "10.5400/jts.2010.15.3.255", "—",
    [("West Aceh oil palm on peat", "Indonesia", "Aceh (West Aceh)", "See paper", "See paper", "—", "Oil palm", "See paper", "See paper",
      "WT depths vs CO2", "See paper", "")])
add("C47", "C45", "rep", "Mires and Peat 2024 (Badas groundwater monitoring)", D + "10.19189/map.2023.cm.sc.2332104", "—",
    [("Badas peat dome, two transects", "Brunei", "Belait (Badas)", "≈4.56", "≈114.34", "Regional approximation", "Degraded peat dome",
      "See paper", "See paper", "Groundwater wells + rain gauge", 1, "2 transects")], "Your C12")
add("C48", "C46", "rep", "Hoyt et al. 2019, GCB", D + "10.1111/gcb.14702", "—",
    [("Damit dome", "Brunei", "Belait", 4.405, 114.363, "Apers et al. 2022 Table B1", "Undrained, previously logged PSF", "2012",
      "See paper", "WT record", 1, "")], "Your A2 (A2 already uses the Damit coordinates)")
add("C49", "C47", "req", "Busman et al. 2023, STOTEN", D + "10.1016/j.scitotenv.2022.159973", "—",
    [("Maludam NP, three forest types", "Malaysia", "Sarawak (Betong)", "See paper", "See paper", "—", "Mixed PSF, Alan batu, Alan bunga",
      "8 years (see paper)", "Monthly", "WT with soil GHG", 3, "")], "Possibly the same plots as your C7")
add("C50", "C48", None, "Imran et al. 2022, Environ. Res. Commun.", D + "10.1088/2515-7620/ac6295", "—",
    [("Sarawak PSF, three sites", "Malaysia", "Sarawak", "See paper", "See paper", "—", "Undrained PSF", "2011-2020", "Monthly",
      "Monthly WT + peat surface level", 3, "")])
add("C51", "C49", "req", "Wong et al. 2018, AFM; Kiew et al. 2020, AFM; Ishikura et al. 2018, AGEE; Ishikura et al. 2019, Ecosystems; Sangok et al. 2017, STOTEN",
    D + "10.1016/j.agrformet.2018.03.025 ; " + D + "10.1016/j.agrformet.2020.108189 ; " + D + "10.1016/j.agee.2017.11.025 ; " + D + "10.1007/s10021-019-00376-8 ; " + D + "10.1016/j.scitotenv.2017.02.165",
    "—",
    [("PSF EC tower (CH4)", "Malaysia", "Sarawak", "See paper", "See paper", "—", "Peat swamp forest", "2014/02-2015/07", "Half-hourly",
      "EC + GWL", 1, ""),
     ("Oil palm EC tower", "Malaysia", "Sarawak", "See paper", "See paper", "—", "Oil palm (7–10 years)", "2011-2014", "Half-hourly", "EC + GWL", 1, ""),
     ("Oil palm automated chambers", "Malaysia", "Sarawak", "See paper", "See paper", "—", "Oil palm", "2 years (see paper)", "Continuous",
      "Automated chambers + GWL", 1, ""),
     ("Undrained PSF automated chambers", "Malaysia", "Sarawak", "See paper", "See paper", "—", "Undrained PSF", "2 years (see paper)",
      "Continuous", "Automated chambers + GWL", 1, ""),
     ("Field incubation, forest peat soils", "Malaysia", "Sarawak", "See paper", "See paper", "—", "PSF peat moved to oil palm",
      "3 years (see paper)", "Monthly", "GWT with CO2/CH4", 3, "")],
    "PSF EC tower: same period as your A6 (likely MY-MLM); data held by the Sarawak government")
add("C52", "C50", None, "Cook et al. 2018, BG", D + "10.5194/bg-15-7435-2018", "—",
    [("Sebungan, Sabaju 1, Sabaju 3, Sabaju 4 estates", "Malaysia", "Sarawak", "See your A12", "See your A12", "—", "Oil palm",
      "2015/08-2016/08", "See paper", "WT depth per site (mean, min, max, % time >60 cm)", 4, "")],
    "Same estates as your A12; data in Cook (2018) thesis")
add("C53", "C51", None, "Basri et al. 2024 (SSRN preprint)", D + "10.2139/ssrn.4767267", "—",
    [("Bintulu oil palm (young and mature)", "Malaysia", "Sarawak (Bintulu)", "See paper", "See paper", "—", "Oil palm",
      "128 months (see paper)", "Monthly", "With soil respiration", "See paper", ""),
     ("Bintulu peat swamp forest", "Malaysia", "Sarawak (Bintulu)", "See paper", "See paper", "—", "PSF", "128 months (see paper)",
      "Monthly", "", "See paper", "")])
add("C54", "C52", None, "Azizan et al. 2021, Water", D + "10.3390/w13233372", "—",
    [("Recovering forest (RF)", "Malaysia", "See paper", "See paper", "See paper", "—", "Recovering PSF", "2017/07-2018/12", "Biweekly",
      "GHG + continuous environmental variables", 1, ""),
     ("Natural forest (NF)", "Malaysia", "See paper", "See paper", "See paper", "—", "Natural PSF", "2017/07-2018/12", "Biweekly", "", 1, ""),
     ("Drained oil palm", "Malaysia", "See paper", "See paper", "See paper", "—", "Oil palm", "2017/07-2018/12", "Biweekly", "", 1, "")])
add("C55", "C53", None, "Katimon et al. 2012, J. Teknologi; Shamsuddin et al. 2021, JWARP",
    D + "10.11113/jt.v38.487 ; " + D + "10.4236/jwarp.2021.1312052", "—",
    [("Parit Madirono peat catchment", "Malaysia", "Johor", "See paper", "See paper", "—", "Drained peat catchment", "See paper",
      "See paper", "WT records", 1, ""),
     ("Ayer Hitam North Forest Reserve, Muar", "Malaysia", "Johor", "See paper", "See paper", "—", "PSF", "See paper", "See paper", "", 1, "")])
add("C56", "C54", None, "Nagano et al. 2013, Mires and Peat", D + "10.19189/001c.128481", "—",
    [("Bacho peatland", "Thailand", "Narathiwat", "See paper", "See paper", "—", "Degraded PSF + conservation zone",
      ">20 years (monthly subsidence)", "Monthly", "WT during soil respiration measurements", 5, "")])
add("C57", "C55", None, "Mires and Peat 2021 (Leyte Sab-a land-use conversion)", D + "10.19189/map.2021.bg.sta.2287", "—",
    [("Leyte Sab-a Basin Peatland", "Philippines", "Leyte", "See paper", "See paper", "—", "PSF vs cultivated", "See paper",
      "Spot", "Spot WT / moisture", "See paper", "")])
add("C58", "C56", None, "Asmat Regency peat survey (ResearchGate 379805559); Apers et al. 2022 (SIPALAGA station)",
    "https://www.researchgate.net/publication/379805559 ; " + D + "10.1029/2021ms002784", "—",
    [("Asmat Regency peat survey", "Indonesia", "Papua (Asmat)", "See paper", "See paper", "—", "Various", "See paper", "Spot",
      "Spot GWL survey", "See paper", ""),
     ("Sumber Mulya SIPALAGA station", "Indonesia", "Papua (Merauke)", -8.205, 140.216, "Apers et al. 2022 Table B1", "Drained",
      "2019-2020", "See source", "SIPALAGA automatic station", 1, "")])
add("C59", "new", "rep", "Apers et al. 2022, JAMES (unpublished sites)", D + "10.1029/2021ms002784", "—",
    [("Raja Musa (IN_N_Selangor)", "Malaysia", "Selangor (North Selangor)", 3.4256, 101.3067, "Apers et al. 2022 Table B1",
      "Drained smallholder agriculture (oil palm, shallow peat)", "2018-2019", "See source", "Unpublished water level record", 1, ""),
     ("Hampangen (IN_Palangkaraya)", "Indonesia", "Central Kalimantan", -1.92, 113.5787, "Apers et al. 2022 Table B1",
      "Drained smallholder agriculture (oil palm, rubber)", "2018-2019", "See source", "", 1, ""),
     ("Teluk Empening (IN_Pontianak)", "Indonesia", "West Kalimantan", -0.3807, 109.5914, "Apers et al. 2022 Table B1",
      "Drained smallholder agriculture (ginger, rubber)", "2018-2019", "See source", "", 1, "")],
    "Coordinates fall inside your A11 (SUSTAINPEAT) regions")

# ---- D: compilations and products (no fill unless replicate)
add("D9", "D2", "rep", "Apers et al. 2022, JAMES", D + "10.1029/2021ms002784", D + "10.5281/zenodo.6011689",
    [("PEATCLSM_Trop evaluation set (87 SE Asian water-level sites)", "Southeast Asia", "Regional", "Per site", "Per site",
      "Apers et al. 2022 Table B1", "All", "2000-2020", "Daily / sub-daily", "Site list: see B5, B13, C15-C18, C48, C59", 87, "")], "Your D2")
add("D10", "D3", "rep", "Koupaei-Abyazani et al. 2024, JGR-Biogeosciences", D + "10.1029/2024jg008116", "—",
    [("OPTRAM satellite WT estimates", "Malaysia, Indonesia, Peru", "Sarawak; Central Kalimantan", "See paper", "See paper", "—",
      "All", "See paper", "Landsat revisit", "Validated against in situ WT", "See paper", "")], "Your D3")
add("D11", "D4", None, "Vernimmen et al. 2020, Water; Dadap et al. 2021, AGU Advances",
    D + "10.3390/w12051486 ; " + D + "10.1029/2020av000321", "—",
    [("Canal water-table depth from airborne LiDAR", "Indonesia", "Eastern Sumatra; West Kalimantan", "Grid", "Grid",
      "Grid resolution 100 m", "Drained peat", "LiDAR dates", "Per survey", "Canal WTD from DTM", "—", ""),
     ("Drainage canal map of SE Asian peatlands", "Southeast Asia", "Regional", "Grid", "Grid", "—", "Drained peat", "See paper",
      "—", "Canal mapping (deep learning)", "—", "")])
add("D12", "D5", None, "Hikouei et al. 2023, STOTEN", D + "10.1016/j.scitotenv.2022.159701", "—",
    [("Machine-learning GWL maps of a drained dome (KFCP data)", "Indonesia", "Central Kalimantan", "Grid", "Grid", "—",
      "Drained / burnt dome", "2010-2012", "See paper", "Modelled GWL", "—", "")])
add("D13", "D6", None, "Mahdiyasa (Zenodo dataset)", "—", D + "10.5281/zenodo.17907349",
    [("Kalimantan peatland fire-risk points with water-table height", "Indonesia", "Kalimantan", "Grid", "Grid",
      "dataset", "All", "Static", "—", "Water table height per point (derived)", "—", "")])
add("D14", "D7", None, "Couwenberg et al. 2010, GCB; Couwenberg & Hooijer 2013, Mires and Peat; Hergoualc'h & Verchot 2014, MASGC; Carlson et al. 2015, ERL; Prananto et al. 2020, GCB; Swails et al. 2024, Biogeochemistry",
    D + "10.1111/j.1365-2486.2009.02016.x ; " + D + "10.19189/001c.128487 ; " + D + "10.1007/s11027-013-9511-x ; " + D + "10.1088/1748-9326/10/7/074006 ; " + D + "10.1111/gcb.15147 ; " + D + "10.1007/s10533-023-01110-2",
    D + "10.17528/CIFOR/DATA.00291",
    [("Meta-analyses with site-level WTD", "Southeast Asia", "—", "See supplement", "See supplement", "—", "Multiple land uses",
      "Literature compilation", "Site mean", "From literature", "Many sites", "")], "Includes your D7 (Prananto 2020)")


# ---- October 2026 additions (IDs continue your updated table: A34, C60-C65, D15)
OCT = "2016-2026 (October 2026 search)"
add("A34", "new", None, "Kagawa et al. 2026, Biogeosciences 23", D + "10.5194/bg-23-2119-2026", D + "10.5281/zenodo.19159833",
    [("Bengkalis Island north-west coast, channel gauge WP1 (peat landslide area)", "Indonesia", "Riau (Bengkalis Island)", "≈1.60",
      "≈102.02", "dataset (UTM 48N survey files)", "Coastal peat with peat mass movements", "2014/12/01-2015/01/31", "10 min",
      "HOBO U-20 logger in a PVC well recording the CHANNEL water level behind a weir (m, elevation; not peat WTD)", 1, "1 channel gauge"),
     ("Bengkalis survey line, one-day water-level survey", "Indonesia", "Riau (Bengkalis Island)", "1.59-1.60", "102.02",
      "dataset (UTM 48N)", "Coastal peat", "2013/08/24", "One survey", "Water-level elevations (m) at 30 points", 1, "30 points")],
    "Same island as C35 and the B9 SESAME station at Dompas", OCT, "Zenodo files; full text")
add("C60", "new", None, "Cassiophea et al. 2025, IOP Conf. Ser. Earth Environ. Sci. 1542", D + "10.1088/1755-1315/1542/1/012023", "—",
    [("RePEAT (rehabilitated reference site)", "Indonesia", "Central Kalimantan", "Map only (Fig. 1)", "Map only (Fig. 1)", "—",
      "Rehabilitated peatland, permanently moist", "2024/11-2025/04", "Hourly to daily (IoT)",
      "Pressure transducer in a 2-3 m PVC well (solar, LTE); monthly manual check", 1, ""),
     ("CIMTROP Lab Forest", "Indonesia", "Central Kalimantan", "Map only (Fig. 1)", "Map only (Fig. 1)", "—",
      "Little-disturbed research forest", "2024/11-2025/04", "Hourly to daily (IoT)", "", 1, ""),
     ("Ruslan Canal, Sebangau NP", "Indonesia", "Central Kalimantan", "Map only (Fig. 1)", "Map only (Fig. 1)", "—",
      "Previously drained, canal-blocked", "2024/11-2025/04", "Hourly to daily (IoT)", "", 1, ""),
     ("KHDTK (Tumbang Nusa forest)", "Indonesia", "Central Kalimantan", "Map only (Fig. 1)", "Map only (Fig. 1)", "—",
      "Forested peatland under restoration (WTD -40 to +25 cm)", "2024/11-2025/04", "Hourly to daily (IoT)", "", 1, ""),
     ("KM 16", "Indonesia", "Central Kalimantan", "Map only (Fig. 1)", "Map only (Fig. 1)", "—",
      "Former degraded shrubland, canal-blocked (ponded +20 to +40 cm)", "2024/11-2025/04", "Hourly to daily (IoT)", "", 1, "")],
    "KHDTK = the Tumbang Nusa forest of your C20/C23; co-author of A11", OCT, "full text")
add("C61", "new", None, "Hikouei et al. 2025, Groundwater for Sustainable Development 29", D + "10.1016/j.gsd.2025.101413",
    "Not shared (authors have no permission)",
    [("KFCP dipwell network, ex-Mega Rice Project", "Indonesia", "Central Kalimantan (Kapuas)", "See B6", "See B6", "—",
      "Degraded, drained peat dome", "2011-2019", "Monthly manual", "Dipwells; MODFLOW model with XGBoost analysis of residuals", 1,
      "265 dipwells")], "Extends your B6 to 2019; same lead author as D12", OCT, "abstract (closed)")
add("C62", "new", None, "Yananto et al. 2021, Int. J. Remote Sens. Earth Sci. 18(2)", "https://ejournal.brin.go.id/ijreses/article/view/13780", "—",
    [("SIPALAGA stations near the Rokan River", "Indonesia", "Riau (Rokan)", "See paper", "See paper", "—", "Acacia plantation",
      "See paper", "See paper", "SIPALAGA GWL regressed on Sentinel-1 VV backscatter (r = -0.648)", "See paper", "")],
    "SIPALAGA stations (your B1, B13)", OCT, "article page")
add("C63", "new", None, "Sutikno et al. 2019, MATEC Web Conf. 276", D + "10.1051/matecconf/201927606003", "—",
    [("Canal block transect", "Indonesia", "Riau (Tebing Tinggi Island)", "See paper", "See paper", "—", "Drained peat with a canal block",
      "See paper", "See paper", "Dipwells at 20, 70, 120, 170 and 220 m from the canal", 1, "5 dipwells")],
    "Same group and island as your C4 and C32", OCT, "abstract (publisher blocked)")
add("C64", "new", None, "Putra et al. 2024, Journal of Tropical Silviculture 15", D + "10.29244/j-siltrop.15.01.65-69", "—",
    [("Tangkit Baru village", "Indonesia", "Jambi (Muaro Jambi, Sungai Gelam)", "Map only", "Map only", "—",
      "Drained pineapple farms (11 canals, 1-2 m deep)", "2023/09-2023/11", "IoT logger (interval not stated) + manual",
      "Submersible pressure sensor (ESP32, Telkom IoT) vs measuring stick in the same dipwell", 1, "1 dipwell"),
     ("Pematang Rahim village", "Indonesia", "Jambi (Tanjung Jabung Timur, Mendahara Ulu)", "Map only", "Map only", "—",
      "Peat next to the Sungai Buluh protection forest; oil palm", "2023/09-2023/11", "IoT logger (interval not stated) + manual", "", 1,
      "1 dipwell")], "Pematang Rahim borders the forest of A17", OCT, "full text")
add("C65", "new", None, "Sekarano et al. 2026, SSRN preprint", D + "10.2139/ssrn.7439436", "—",
    [("26 GWL stations, Pulang Pisau and Palangka Raya", "Indonesia", "Central Kalimantan", "See paper", "See paper", "—",
      "Degraded and restored peat", "2018/12-2022/12", "21,954 observations (about daily, with gaps)",
      "Station GWL vs Sentinel-2 bands and moisture indices; LSTM (R2 0.89, RMSE 0.13 m)", 26, "")],
    "Station network not named; probably SIPALAGA (your B1)", OCT, "abstract")
add("D15", "new", None, "Irfan et al. (submitted to Science of the Total Environment)", "—", D + "10.5281/zenodo.23008826",
    [("South Sumatra peat units (KHG): modelled GWL maps", "Indonesia", "South Sumatra", "Grid", "Grid", "1 km grid (EPSG:32748)",
      "All peat", "2019/04/01-13 (wet); 2019/11/13-25 (dry)", "Two composites",
      "Ridge regression of Sentinel-1, GPM and SMAP, calibrated on SIPALAGA GWL (field data not included)", "—", "")],
    "Same group as your B9 and C42", OCT, "Zenodo files and README")

UPDATES = [
    ("A3 / A4", "B4: Sulaiman et al. 2023 report daily GWL at the UF location from 1993/09 to 2019/12."),
    ("A2", "Apers et al. 2022 list two Hoyt et al. 2019 records: Damit dome 2012 (4.405, 114.363) and Mendaram dome 2013-2014 "
           "(4.3599, 114.3522). Your A2 uses the Damit coordinates for the whole 2012-2015 record; Period 1 and Period 2 may be at different domes."),
    ("A4 / A6", "A15: the daily WT behind your A4 rows 9-10 (Maludam 2011-2014; Naman oil palm 2018-2019) is public on figshare. "
                "A16: half-hourly WT at MY-MLM for Nov-Dec 2013 on Zenodo."),
    ("A8", "C37: Deshmukh et al. 2020 give the same Acacia tower position (0°30'57\"N, 102°02'E), so 0.51589, 102.03641 is correct."),
    ("A11", "C59: the three 'unpublished' Apers sites (Raja Musa, Hampangen, Teluk Empening) lie inside your SUSTAINPEAT regions."),
    ("A12", "C52: Cook et al. 2018 measured WT at Sebungan and Sabaju 1/3/4 (2015/08-2016/08)."),
    ("A14", "A18: Swails et al. 2021 (CH4/N2O, CIFOR DATA.00201) probably uses the same six sites."),
    ("B1", "B13: 59 SIPALAGA station codes with coordinates; Putra et al. 2025 used 39 Riau stations (2018/10-2020/12, daily)."),
    ("Gaps", "Thailand (B10, C56), Vietnam (B11) and Papua (C58) now have leads; Philippines only a spot survey (C57); "
             "Sabah (Klias) still has only geophysical WT estimates."),
]

E.sort(key=lambda e: (e["new"][0], int(e["new"][1:])))

# ---------------------------------------------------------------- corrections from reading the papers and files
COLS = ["ID", "Source publication(s)", "Publication link(s)", "Site / dataset", "Country", "Province / state", "Latitude",
        "Longitude", "Source", "Land use / condition", "Period", "Temporal resolution", "Method / accuracy", "Sites", "Wells",
        "Data link"]
ROW_COLS = COLS[3:15]                       # the 12 per-site cells, in the order of the row tuples
ENTRY_COLS = {"Source publication(s)": "pubs", "Publication link(s)": "links", "Data link": "data"}
LABELS = {"can use": "can", "replicate": "rep", "on request": "req", "not accessible": "na"}
NAMES = {v: k for k, v in LABELS.items()}
MATCH = ("Site / dataset", "Province / state", "Land use / condition", "Method / accuracy")  # where a fix's row text is looked for
KEY_COLS = (11, 12, 14)                     # Period, Temporal resolution, Sites: dark red in the template


def colour(value):
    """Legend key of a 'Colour' fix such as 'replicate (= your A28)'."""
    for label, key in LABELS.items():
        if str(value).lower().startswith(label):
            return key
    raise ValueError("a Colour fix must start with a legend label: %r" % (value,))


def same(a, b):
    return str(a if a is not None else "").strip() == str(b if b is not None else "").strip()


def pick_rows(texts, spec):
    """Indices of the rows a fix applies to. texts[k] = (site cell, province + land use + method cells) of row k;
    spec '' = first row, 'text' = first row containing text, '*text' = every row containing text, '*' = every row.
    The site cells are searched first; the other cells only if no site matches."""
    if not spec:
        return [0] if texts else []
    every, needle = spec.startswith("*"), spec.lstrip("*").lower()
    for part in (0, 1):
        hits = [k for k, t in enumerate(texts) if needle in str(t[part] or "").lower()]
        if hits:
            return hits if every else hits[:1]
    return []


def fixes_for(scope, fixes=TABLE_FIXES):
    """The fixes that apply to 'table' (your table) or 'expanded' (EXPANDED_2006-2026.xlsx), as 6-tuples."""
    out = []
    for fx in fixes:
        fx = tuple(fx) + ("",) * (6 - len(fx))
        if fx[5] in ("", scope):
            out.append(fx)
    return out


def apply_fixes(entries, fixes=None):
    """Write the fixes into the entries; return the cells that changed as (ID, site, column, old, new, source)."""
    by_id = {e["new"]: e for e in entries}
    changes = []
    for i, spec, col, new, src, _ in fixes_for("expanded") if fixes is None else fixes:
        e = by_id.get(i)
        if e is None:
            continue                        # a row of your table that EXPANDED does not have
        if col == "New row":
            if not any(same(r[0], new[0]) for r in e["rows"]):
                e["rows"].append(tuple(new))
                changes.append((i, new[0], col, "", " | ".join(str(x) for x in new), src))
                e.setdefault("marks", []).append((len(e["rows"]) - 1, "Site / dataset", None, src))
            continue
        texts = [(r[0], " ".join(str(r[ROW_COLS.index(c)]) for c in MATCH[1:])) for r in e["rows"]]
        ks = pick_rows(texts, spec)
        if col in ("Colour",) or col in ENTRY_COLS:
            ks = ks[:1]                     # entry-level cells: once
        for k in ks:
            if col == "Colour":
                old = NAMES.get(e["cat"], "no fill")
                if old == NAMES[colour(new)]:
                    continue
                e["cat"] = colour(new)
            elif col in ENTRY_COLS:
                old = e[ENTRY_COLS[col]]
                e[ENTRY_COLS[col]] = new
            else:
                row = list(e["rows"][k])
                j = ROW_COLS.index(col)
                old, row[j] = row[j], new
                e["rows"][k] = tuple(row)
            if not same(old, new):
                changes.append((i, e["rows"][k][0], col, old, new, src))
                e.setdefault("marks", []).append((0 if col in ENTRY_COLS or col == "Colour" else k,
                                                  "ID" if col == "Colour" else col, old, src))
    changes.sort(key=lambda c: (c[0][0], int(c[0][1:])))  # by ID; within an ID, in the order of TABLE_FIXES
    return changes


# ---------------------------------------------------------------- write
def load_template(path):
    """The user's table as a template: the sheet whose first row is the column header and which has the colour legend
    (other sheets are removed), the legend fills (found by their labels) and the fonts of the first data row."""
    wb = openpyxl.load_workbook(path)
    best = None
    for ws in wb.worksheets:
        if [str(ws.cell(1, c).value or "").strip() for c in range(1, len(COLS) + 1)] != COLS:
            continue
        legend = {}
        for r in range(2, ws.max_row + 1):
            label = str(ws.cell(r, 2).value or "").strip().lower()
            if label in LABELS and ws.cell(r, 1).value is None and ws.cell(r, 1).fill.fill_type:
                legend[LABELS[label]] = (r, ws.cell(r, 2).value)
        if best is None or len(legend) > len(best[1]):
            best = (ws, legend)
    if best is None or len(best[1]) < len(LABELS):
        sys.exit("%s: no sheet with the SEA_peatland_WTD_datasets header row and colour legend" % path)
    ws, legend = best
    for other in list(wb.worksheets):
        if other is not ws:
            wb.remove(other)
    link = copy.copy(ws["P2"].font)
    if not link.u:
        link = openpyxl.styles.Font(name=ws["D2"].font.name, sz=ws["D2"].font.sz, color="FF0000FF", u="single")
    st = dict(fills={k: copy.copy(ws.cell(r, 1).fill) for k, (r, _) in legend.items()},
              legend=[(k, legend[k][1]) for k in ("can", "rep", "req", "na")],
              head=copy.copy(ws["A1"].font), base=copy.copy(ws["D2"].font), key=copy.copy(ws["K2"].font), link=link,
              align=copy.copy(ws["D2"].alignment))
    ws.delete_rows(2, ws.max_row)
    for k in list(ws.row_dimensions):       # delete_rows does not shift row heights
        if k > 1:
            del ws.row_dimensions[k]
    return wb, ws, st


def write_rows(ws, st, entries, r=2):
    """Write entries in the template layout from row r: ID, publications and data link on the first row of each entry;
    the entry's colour on its filled cells, as in the template. Returns the next free row."""
    for e in entries:
        e["xlrow"] = r
        for i, row in enumerate(e["rows"]):
            vals = [e["new"] if i == 0 else None, e["pubs"] if i == 0 else None, e["links"] if i == 0 else None]
            vals += list(row[:10]) + [row[10] if row[10] != "" else None, row[11] or None, e["data"] if i == 0 else None]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c, v)
                cell.font = copy.copy(st["key"] if c in KEY_COLS else st["base"])
                cell.alignment = copy.copy(st["align"])
                if e["cat"] and v is not None:
                    cell.fill = copy.copy(st["fills"][e["cat"]])
                if c in (3, 16) and isinstance(v, str) and v.startswith("http"):
                    cell.hyperlink = v.split(" ")[0]
                    cell.font = copy.copy(st["link"])
            r += 1
    return r


def mark_changes(ws, entries):
    """A comment on every cell that a fix changed: the source, and the value it replaced."""
    for e in entries:
        for k, col, old, src in e.get("marks", []):
            if col == "Site / dataset" and old is None:
                text = "Row added after reading the papers (%s)." % src
            else:
                text = "Changed after reading the papers (%s). Was: %s" % (src, "(empty)" if old in (None, "") else old)
            c = ws.cell(e["xlrow"] + k, COLS.index(col) + 1)
            c.comment = openpyxl.comments.Comment(text, "literature check 2026-10", width=320, height=110)


def write_legend(ws, st, r):
    for k, label in st["legend"]:
        ws.cell(r, 1).fill = copy.copy(st["fills"][k])
        ws.cell(r, 2, label).font = copy.copy(st["base"])
        r += 1


def plain_sheet(wb, st, title, header, rows, widths):
    sh = wb.create_sheet(title)
    sh.append(header)
    for c in sh[1]:
        c.font = copy.copy(st["head"])
    for row in rows:
        sh.append(["" if v is None else v for v in row])
    for row in sh.iter_rows(min_row=2):
        for c in row:
            c.font = copy.copy(st["base"])
            c.alignment = openpyxl.styles.Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        sh.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    sh.freeze_panes = "A2"
    return sh


def main(template, out):
    changes = apply_fixes(E)
    wb, ws, st = load_template(template)
    write_legend(ws, st, write_rows(ws, st, E) + 3)
    mark_changes(ws, E)

    # ID mapping (keeps what the template layout has no column for)
    NOTE_FIX = [("(see C45)", "(see C47)"), ("Hikouei et al. 2023, D6)", "Hikouei et al. 2023, D12)"),
                ("SATREPS IJ-1 (B4)", "SATREPS IJ-1 (B5)"), ("in INDEX row C10-11", "in your C10/C11")]
    labels = dict(st["legend"])
    rows = []
    for e in E:
        o = OLD.get(e["old"])
        note = o[12] if o and o[12] else ""
        for a, b in NOTE_FIX:               # notes were written with EXPANDED / INDEX IDs
            note = note.replace(a, b)
        rows.append([e["new"], e["old"], labels.get(e["cat"], "—") if e["cat"] else "—", e["dup"] or "—",
                     e["window"] or (o[10] if o else "2016-2026"), e["via"] or (o[11] if o else "Apers 2022 Table B1"), note])
    plain_sheet(wb, st, "ID mapping", ["New ID", "EXPANDED ID", "Colour", "Duplicate of / related to (your table)",
                                       "Publication window", "Verified via", "Notes"], rows, [8, 12, 14, 48, 30, 26, 70])
    plain_sheet(wb, st, "Updates to your table", ["Your entry", "Update"], [list(u) for u in UPDATES], [10, 130])
    plain_sheet(wb, st, "Changes after reading papers", ["ID", "Row (site)", "Column", "Old value", "New value", "Source"],
                [list(c) for c in changes], [7, 40, 20, 40, 60, 40])
    wb.active = 0
    wb.properties.lastModifiedBy = None
    wb.save(out)
    print("entries %d, rows %d, changed cells %d -> %s" % (len(E), sum(len(e["rows"]) for e in E), len(changes), out))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "EXPANDED_2006-2026.xlsx"))
